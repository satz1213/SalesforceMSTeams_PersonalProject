"""Generate Salesforce-Teams Case Swarm architecture Word document."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUTPUT = Path(__file__).resolve().parent / "Salesforce-Teams-Case-Swarm-Architecture.docx"


def set_document_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def add_title_page(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Salesforce ↔ Microsoft Teams\nCase Swarm Integration")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 112, 210)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Architecture & Component Reference").font.size = Pt(16)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Document generated: {date.today():%B %d, %Y}\n")
    meta.add_run("Project: SFMSTeamsInOutProject")

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * level)


def add_numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph()


def add_code_block(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.25)


def build_document() -> Document:
    doc = Document()
    set_document_styles(doc)
    add_title_page(doc)

    # 1. Overview
    add_heading(doc, "1. Executive Overview")
    doc.add_paragraph(
        "This solution enables Salesforce Service Cloud agents to start a Microsoft Teams "
        "\"swarm\" directly from a Case record. A swarm brings subject-matter experts (SMEs) "
        "into a shared Teams workspace—either a Team (M365 group) or a group Chat—so they can "
        "collaborate on resolving the case without leaving their existing tools."
    )
    doc.add_paragraph(
        "For Chat swarms, an optional Teams bot posts an interactive Adaptive Card that lets "
        "participants view and edit Case fields (Status, Priority, Owner, Subject, Description) "
        "inline inside Teams, with changes written back to Salesforce."
    )
    doc.add_paragraph("Key design principle: all Microsoft API calls run server-side in Apex. "
                      "The Lightning Web Component (LWC) never calls Microsoft Graph or Bot Framework directly.")

    # 2. High-level architecture
    add_heading(doc, "2. High-Level Architecture")
    doc.add_paragraph(
        "The integration spans three platforms: Salesforce (business logic & data), "
        "Microsoft 365 / Teams (collaboration), and Azure (bot registration & HTTP relay)."
    )
    add_code_block(doc, """
┌─────────────────────────────────────────────────────────────────────────────┐
│                              SALESFORCE ORG                                  │
│  ┌──────────────┐    ┌─────────────────────┐    ┌────────────────────────┐  │
│  │  caseSwarm   │───▶│ TeamsSwarmController │───▶│  TeamsSwarmService     │  │
│  │  (LWC)       │    │ (Apex @AuraEnabled)  │    │  (Graph API callouts)  │  │
│  └──────────────┘    └─────────────────────┘    └───────────┬────────────┘  │
│                                                               │              │
│  ┌──────────────────────────────────────────────────────────┐ │              │
│  │ TeamsBotMessagingResource (REST /teamsbot/messages)      │◀┼──────────────┤
│  │  ├─ TeamsBotJwtValidator                                 │ │              │
│  │  ├─ TeamsAdaptiveCardBuilder                             │ │              │
│  │  └─ TeamsBotConversationService (Bot Connector API)      │─┼──────────────┤
│  └──────────────────────────────────────────────────────────┘ │              │
│         Named Credentials: MS_Graph, MS_Bot_Framework          │              │
└─────────────────────────────────────────────────────────────────┼──────────────┘
                                                                  │
                    Microsoft Graph API ◀─────────────────────────┘
                    Bot Framework Connector API ◀───────────────────┘
                                  │
┌─────────────────────────────────▼───────────────────────────────────────────┐
│                           MICROSOFT AZURE / M365                               │
│  ┌─────────────────┐   ┌──────────────────┐   ┌───────────────────────────┐ │
│  │ Azure AD App    │   │ Azure Bot Service │   │ Azure Function (relay)    │ │
│  │ Registration    │   │ (messaging ep.)   │──▶│ teamsBotRelay             │ │
│  │ (Graph + Bot)   │   └─────────┬─────────┘   └─────────────┬─────────────┘ │
│  └─────────────────┘             │                           │               │
│                                  ▼                           ▼               │
│                         Microsoft Teams              Back to Salesforce REST │
│                         (Team / Chat / Bot)          (OAuth + header rewrite)│
└───────────────────────────────────────────────────────────────────────────────┘
""")

    # 3. Why not direct browser calls
    add_heading(doc, "3. Why Server-Side Integration?")
    add_bullet(doc, "Authentication: Microsoft Graph requires OAuth bearer tokens. Client secrets and app-only tokens cannot safely live in browser code (LWC).")
    add_bullet(doc, "CORS: graph.microsoft.com does not accept arbitrary browser-origin requests for app-only auth flows.")
    add_bullet(doc, "Resulting pattern: LWC → Apex Controller → Apex Service → Named Credential → Microsoft APIs.")

    # 4. Components
    add_heading(doc, "4. Component Inventory")

    add_heading(doc, "4.1 Salesforce – User Interface", level=2)
    add_table(doc,
        ["Component", "Type", "Purpose"],
        [
            ["caseSwarm", "Lightning Web Component (LWC)", "Case record page UI: SME search, swarm type selection (Team/Chat), Start swarm / Add members, Open in Teams link, provisioning status polling."],
        ])

    add_heading(doc, "4.2 Salesforce – Apex Controllers & Services", level=2)
    add_table(doc,
        ["Component", "Purpose", "Key Connections"],
        [
            ["TeamsSwarmController", "LWC entry point. searchUsers(), startSwarm(), addMembersToSwarm().", "Calls TeamsSwarmService; enqueues async jobs."],
            ["TeamsSwarmService", "Microsoft Graph API: create M365 group, enable Team, add members, create Chat, install bot app.", "Named Credential: MS_Graph (OAuth client credentials)."],
            ["TeamsSwarmProvisionQueueable", "Retries PUT /groups/{id}/team until Graph finishes propagating new group.", "MS_Graph; updates Case Swarm_Status__c."],
            ["TeamsSwarmProvisionScheduler", "Schedules first provisioning retry attempt.", "Enqueues TeamsSwarmProvisionQueueable."],
            ["TeamsSwarmBotInstallQueueable", "Installs bot into Chat after swarm creation (async).", "TeamsSwarmService.installBotApp() via Graph."],
            ["TeamsBotMessagingResource", "REST endpoint /services/apexrest/teamsbot/messages. Handles Bot Framework activities.", "Validates JWT; updates Cases; posts cards."],
            ["TeamsBotJwtValidator", "Validates Bot Framework RS256 JWT against login.botframework.com JWKS.", "Remote Site: Bot_Framework_Keys."],
            ["TeamsAdaptiveCardBuilder", "Builds view/edit Adaptive Card JSON (Action.Execute model).", "Pure JSON builder; no callouts."],
            ["TeamsBotConversationService", "Posts Adaptive Cards via Bot Framework Connector API v3.", "Named Credential: MS_Bot_Framework."],
            ["TeamsBotInstalledCardQueueable", "Posts initial case card after bot installation (async).", "TeamsBotConversationService."],
        ])

    add_heading(doc, "4.3 Salesforce – Configuration & Data", level=2)
    add_table(doc,
        ["Resource", "Purpose"],
        [
            ["MS_Graph / MS_Graph_Cred", "Named Credential + External Credential for Microsoft Graph (client credentials)."],
            ["MS_Bot_Framework / MS_Bot_Framework_Cred", "Named Credential + External Credential for Bot Framework Connector API."],
            ["Teams_Bot_Config__mdt (Default)", "Stores Microsoft App Id and Teams App Catalog Id for bot install."],
            ["Teams_Swarm_User permission set", "Grants access to swarm Apex classes for agents."],
            ["Case custom fields", "Swarm_Team_Id__c, Swarm_Group_Id__c, Swarm_Team_Url__c, Swarm_Status__c, Swarm_Type__c, Swarm_Bot_Service_Url__c, Swarm_Bot_Install_Error__c."],
            ["User custom fields", "Azure_AD_Email_Id__c (required), Azure_AD_Object_Id__c (optional reference)."],
        ])

    add_heading(doc, "4.4 Microsoft Azure", level=2)
    add_table(doc,
        ["Component", "Purpose", "Connection Details"],
        [
            ["Azure AD App Registration", "Service identity for Graph API and optionally the Teams bot.", "Application permissions: Group.ReadWrite.All, Team.Create, Directory.Read.All, Chat.ReadWrite.All, TeamsAppInstallation.ReadWriteForChat.All. Client secret for OAuth."],
            ["Azure Bot Service", "Registers the bot; receives Teams activities; forwards to messaging endpoint.", "Messaging endpoint → Azure Function relay URL (NOT Salesforce directly)."],
            ["Azure Function: teamsBotRelay", "HTTP proxy: rewrites Authorization header, authenticates to Salesforce, forwards body.", "App settings: SALESFORCE_MESSAGING_ENDPOINT, SALESFORCE_LOGIN_URL, SALESFORCE_CLIENT_ID, SALESFORCE_CLIENT_SECRET."],
            ["Teams App Catalog entry", "Custom Teams app (teamsapp/manifest.json) uploaded to tenant.", "Enables bot install into group chats via Graph installedApps API."],
        ])

    # 5. Authentication flows
    add_heading(doc, "5. Authentication & Connection Establishment")

    add_heading(doc, "5.1 Salesforce → Microsoft Graph (Swarm Creation)", level=2)
    add_numbered(doc, "Agent clicks Start swarm in LWC on a Case record.")
    add_numbered(doc, "LWC calls TeamsSwarmController.startSwarm() via @AuraEnabled Apex.")
    add_numbered(doc, "Apex calls TeamsSwarmService methods using HttpRequest with endpoint callout:MS_Graph.")
    add_numbered(doc, "Salesforce Named Credential MS_Graph automatically obtains an OAuth 2.0 access token from https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token using client credentials (MS_Graph_Cred External Credential).")
    add_numbered(doc, "Scope: https://graph.microsoft.com/.default. Token attached as Authorization: Bearer on every Graph call.")
    add_numbered(doc, "Graph creates M365 group / Chat, adds members identified by User.Azure_AD_Email_Id__c (UPN/email).")

    add_heading(doc, "5.2 Salesforce → Bot Framework Connector (Posting Cards)", level=2)
    add_numbered(doc, "When bot is installed, TeamsBotMessagingResource captures conversation serviceUrl and stores it on Case.Swarm_Bot_Service_Url__c.")
    add_numbered(doc, "TeamsBotConversationService posts Adaptive Cards to {serviceUrl}/v3/conversations/{id}/activities.")
    add_numbered(doc, "Authenticated via MS_Bot_Framework Named Credential (separate OAuth client credentials, scope https://api.botframework.com/.default).")

    add_heading(doc, "5.3 Microsoft Bot Framework → Salesforce (Inbound Activities)", level=2)
    doc.add_paragraph(
        "This is the most complex connection and requires the Azure Function relay."
    )
    add_numbered(doc, "Teams user action (bot install, Edit Case, Save) triggers Bot Framework to POST an Activity JSON to the Azure Bot messaging endpoint.")
    add_numbered(doc, "Azure Bot Service forwards the request to the Azure Function teamsBotRelay.")
    add_numbered(doc, "Relay reads Bot Framework JWT from Authorization header.")
    add_numbered(doc, "Relay obtains Salesforce OAuth token via client credentials (Connected App) from SALESFORCE_LOGIN_URL/services/oauth2/token.")
    add_numbered(doc, "Relay forwards POST to Salesforce /services/apexrest/teamsbot/messages with: Authorization: Bearer {Salesforce token}; X-Bot-Framework-Authorization: {Bot JWT}; body unchanged.")
    add_numbered(doc, "TeamsBotMessagingResource validates X-Bot-Framework-Authorization via TeamsBotJwtValidator (signature, issuer, audience = bot App Id, expiry).")
    add_numbered(doc, "Resource processes activity: conversationUpdate/installationUpdate → post card; invoke adaptiveCard/action → edit/save Case and return updated card JSON synchronously.")

    add_heading(doc, "5.4 Why the Relay Exists (Not Power Automate)", level=2)
    add_bullet(doc, "Salesforce REST dispatcher intercepts Authorization on /services/apexrest/* and validates it as a Salesforce session before Apex runs. Bot Framework JWT in that header causes platform 401.")
    add_bullet(doc, "Relay moves Bot JWT to X-Bot-Framework-Authorization and uses its own Salesforce OAuth token in Authorization.")
    add_bullet(doc, "Azure Function chosen over Power Automate: synchronous low-latency pass-through required for Teams invoke responses; minimal plumbing; no business logic drift; standard Azure Bot endpoint pattern.")

    # 6. User flows
    add_heading(doc, "6. End-to-End User Flows")

    add_heading(doc, "6.1 Start Team Swarm", level=2)
    add_numbered(doc, "Agent searches and selects SMEs (users with Azure_AD_Email_Id__c populated).")
    add_numbered(doc, "Selects swarm type: Team.")
    add_numbered(doc, "TeamsSwarmService.createGroup() → POST /v1.0/groups.")
    add_numbered(doc, "Case updated: Swarm_Status__c = Provisioning, Swarm_Type__c = Team, URLs stored.")
    add_numbered(doc, "TeamsSwarmProvisionScheduler enqueues retry job for PUT /groups/{id}/team (Graph propagation delay).")
    add_numbered(doc, "On success: members added via Graph $batch; Swarm_Status__c = Active.")
    add_numbered(doc, "LWC shows Open in Teams button.")

    add_heading(doc, "6.2 Start Chat Swarm (with Interactive Bot Card)", level=2)
    add_numbered(doc, "Agent selects swarm type: Chat.")
    add_numbered(doc, "TeamsSwarmService.createChat() → POST /v1.0/chats with members.")
    add_numbered(doc, "Case updated immediately: Swarm_Status__c = Active, Swarm_Type__c = Chat.")
    add_numbered(doc, "TeamsSwarmBotInstallQueueable (async) calls Graph POST /chats/{id}/installedApps to install bot.")
    add_numbered(doc, "Teams sends conversationUpdate/installationUpdate to bot messaging endpoint → relay → TeamsBotMessagingResource.")
    add_numbered(doc, "Resource stores serviceUrl on Case; posts initial view Adaptive Card via TeamsBotConversationService.")
    add_numbered(doc, "User clicks Edit Case → invoke activity → card swaps to edit form in place.")
    add_numbered(doc, "User clicks Save → Case updated in Salesforce (scoped to conversation id); view card returned.")

    add_heading(doc, "6.3 Add Members to Existing Swarm", level=2)
    add_numbered(doc, "Agent selects additional SMEs on active swarm Case.")
    add_numbered(doc, "TeamsSwarmController.addMembersToSwarm() calls Graph addMembers (Team) or addChatMembers (Chat).")

    # 7. Data model
    add_heading(doc, "7. Identity & Data Model")
    doc.add_paragraph(
        "Users must have Azure_AD_Email_Id__c populated with their Microsoft 365 email/UPN. "
        "This bridges Salesforce User records to Microsoft Graph user references without requiring "
        "Azure AD Object IDs in most paths."
    )
    add_table(doc,
        ["Case Field", "Description"],
        [
            ["Swarm_Team_Id__c", "Graph Team id or Chat id (conversation id for chats)."],
            ["Swarm_Group_Id__c", "M365 group id (Team swarms only)."],
            ["Swarm_Team_Url__c", "Deep link to open swarm in Teams client."],
            ["Swarm_Status__c", "Empty, Provisioning, Active, or Failed."],
            ["Swarm_Type__c", "Team or Chat."],
            ["Swarm_Bot_Service_Url__c", "Bot Framework regional serviceUrl for posting cards."],
            ["Swarm_Bot_Install_Error__c", "Diagnostic message if bot install fails."],
        ])

    # 8. Security
    add_heading(doc, "8. Security Model")
    add_bullet(doc, "Graph & Bot outbound callouts: OAuth 2.0 client credentials stored in Salesforce External Credentials (not in code).")
    add_bullet(doc, "Inbound bot requests: Bot Framework JWT validated in TeamsBotJwtValidator (RS256, issuer, audience, expiry).")
    add_bullet(doc, "Relay Salesforce OAuth proves the integration caller is authorized; Bot JWT proves the request originated from Bot Framework.")
    add_bullet(doc, "Case updates from card actions are scoped server-side: Case.Swarm_Team_Id__c must match the activity conversation id—prevents cross-case tampering via manipulated caseId in card data.")
    add_bullet(doc, "TeamsSwarmController and messaging resource use with sharing / user-mode SOQL where applicable for agent-initiated flows.")

    # 9. Deployment checklist
    add_heading(doc, "9. Deployment Checklist")
    add_table(doc,
        ["Step", "Platform", "Action"],
        [
            ["1", "Azure AD", "Register app; grant Graph application permissions; create client secret."],
            ["2", "Salesforce", "Configure MS_Graph Named Credential + External Credential."],
            ["3", "Salesforce", "Create Case/User custom fields; populate Azure_AD_Email_Id__c."],
            ["4", "Salesforce", "Deploy force-app package; add caseSwarm LWC to Case page."],
            ["5", "Azure", "Create Azure Bot; set messaging endpoint to relay URL."],
            ["6", "Azure", "Deploy teamsBotRelay Function App with Salesforce app settings."],
            ["7", "Salesforce", "Create Connected App for relay OAuth; configure relay credentials."],
            ["8", "Teams", "Upload teamsapp manifest to tenant app catalog; note catalog id."],
            ["9", "Salesforce", "Configure Teams_Bot_Config__mdt, MS_Bot_Framework Named Credential."],
            ["10", "Salesforce", "Expose REST endpoint (Site or internal URL per your deployment)."],
        ])

    # 10. Known limitations
    add_heading(doc, "10. Known Limitations")
    add_bullet(doc, "Interactive Adaptive Card bot works for Chat swarms only—not Team/channel swarms.")
    add_bullet(doc, "Team provisioning is asynchronous; Case shows Provisioning until Graph completes team enablement.")
    add_bullet(doc, "JWKS fetched on every bot request (acceptable at low volume).")
    add_bullet(doc, "Guest User Site approach retired in favor of relay OAuth as real integration user (avoids Secure Guest User record access restrictions on Case).")

    # Footer note
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("— End of document —")
    run.italic = True
    run.font.color.rgb = RGBColor(128, 128, 128)

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
